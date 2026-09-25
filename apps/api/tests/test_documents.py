import io
import zipfile

import pytest
from docx import Document

from app.documents.chunking import TARGET_CHARS, chunk_sections
from app.documents.parsing import DocumentParseError, Section, parse_document
from app.documents.validation import (
    UploadRejected,
    content_problem,
    detect_kind,
    sanitize_filename,
)

from .conftest import README, make_pdf


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("../../etc/passwd.txt", "passwd.txt"),
        ("C:\\Users\\me\\Proposal Final.docx", "Proposal Final.docx"),
        ("notes\x00\x07.md", "notes.md"),
        ("<script>.md", "_script_.md"),
        ("", "document"),
        (None, "document"),
    ],
)
def test_sanitize_filename(raw: str | None, expected: str) -> None:
    assert sanitize_filename(raw) == expected


def test_sanitize_filename_truncates_but_keeps_extension() -> None:
    name = sanitize_filename("a" * 300 + ".pdf")
    assert len(name) <= 120
    assert name.endswith(".pdf")


def test_detect_kind() -> None:
    assert detect_kind("README.MD") == "md"
    assert detect_kind("spec.markdown") == "md"
    with pytest.raises(UploadRejected) as exc:
        detect_kind("setup.exe")
    assert exc.value.status_code == 415


def test_content_problem_checks_real_bytes() -> None:
    assert content_problem(make_pdf(["hi"]), "pdf") is None
    assert "valid PDF" in (content_problem(b"not a pdf", "pdf") or "")
    assert "binary" in (content_problem(b"MZ\x00\x00", "txt") or "")
    assert "valid DOCX" in (content_problem(b"PK\x03\x04garbage", "docx") or "")

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("other.xml", "<x/>")
    assert "valid DOCX" in (content_problem(archive.getvalue(), "docx") or "")


def test_docx_zip_bomb_guard() -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", "<w/>")
        zf.writestr("word/huge.bin", b"\0" * (65 * 1024 * 1024))
    assert "64 MB" in (content_problem(archive.getvalue(), "docx") or "")


def test_parse_markdown_sections_by_heading() -> None:
    sections = parse_document(README, "md")
    headings = [s.heading for s in sections]
    assert headings == ["Notewave", "Realtime sync", "Conflict handling", "Roadmap"]
    assert "optimistic locking" in sections[2].text


def test_parse_markdown_ignores_headings_inside_code_blocks() -> None:
    sections = parse_document(b"# Setup\n```bash\n# not a heading\nnpm i\n```\n", "md")
    assert [s.heading for s in sections] == ["Setup"]


def test_parse_pdf_keeps_page_numbers() -> None:
    sections = parse_document(
        make_pdf(["Architecture overview", "Conflict handling uses versions"]), "pdf"
    )
    assert [s.page for s in sections] == [1, 2]
    assert "Conflict handling" in sections[1].text


def test_parse_pdf_without_text_fails_cleanly() -> None:
    with pytest.raises(DocumentParseError):
        parse_document(make_pdf([""]), "pdf")


def test_parse_docx_headings_and_tables() -> None:
    document = Document()
    document.add_heading("Architecture", level=1)
    document.add_paragraph("Audio streams to AssemblyAI over WebSocket.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Database"
    table.rows[0].cells[1].text = "Supabase"
    buffer = io.BytesIO()
    document.save(buffer)

    sections = parse_document(buffer.getvalue(), "docx")
    assert sections[0].heading == "Architecture"
    assert "WebSocket" in sections[0].text
    assert sections[-1].text == "Database | Supabase"


def test_parse_text_falls_back_from_utf8() -> None:
    assert parse_document("café".encode("cp1252"), "txt")[0].text == "café"


def test_chunking_splits_long_sections_with_overlap() -> None:
    paragraph = "Optimistic locking protects every note from lost updates. " * 12
    sections = [Section(text="\n\n".join([paragraph] * 4), page=3, heading="Conflicts")]
    chunks = chunk_sections(sections, "doc_1", "architecture.pdf")

    assert len(chunks) > 1
    assert all(len(c.content) <= TARGET_CHARS * 2 for c in chunks)
    assert [c.id for c in chunks] == [f"doc_1:{i}" for i in range(len(chunks))]
    assert chunks[0].location == "p. 3 · Conflicts"
    # The next chunk starts with the tail of the previous one.
    assert chunks[1].content.split("\n\n")[0] in chunks[0].content
