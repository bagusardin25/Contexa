"""Text extraction for uploaded documents. Content is only read, never executed."""

import io
import re
from dataclasses import dataclass

from app.models.session import DocumentKind

MAX_PDF_PAGES = 300
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None
    heading: str | None = None


class DocumentParseError(Exception):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Older Windows editors save as cp1252; keep the document usable.
        return data.decode("cp1252", errors="replace")


def _parse_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentParseError("This PDF is password-protected.")
        sections = []
        for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
            text = _clean(page.extract_text() or "")
            if text:
                sections.append(Section(text=text, page=number))
    except DocumentParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise DocumentParseError("This PDF couldn't be read.") from exc

    if not sections:
        raise DocumentParseError("No text found in this PDF. Scanned PDFs aren't supported yet.")
    return sections


def _parse_docx(data: bytes) -> list[Section]:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise DocumentParseError("This DOCX document couldn't be read.") from exc

    sections: list[Section] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        text = _clean("\n\n".join(buffer))
        if text:
            sections.append(Section(text=text, heading=heading))
        buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            flush()
            heading = text
        else:
            buffer.append(text)
    flush()

    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = _clean("\n".join(row for row in rows if row.strip(" |")))
        if text:
            sections.append(Section(text=text, heading="Table"))

    if not sections:
        raise DocumentParseError("This DOCX document has no text.")
    return sections


def _parse_markdown(data: bytes) -> list[Section]:
    sections: list[Section] = []
    heading: str | None = None
    buffer: list[str] = []
    in_code = False

    def flush() -> None:
        text = _clean("\n".join(buffer))
        if text:
            sections.append(Section(text=text, heading=heading))
        buffer.clear()

    for line in _decode_text(data).splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
        match = None if in_code else _HEADING_RE.match(line)
        if match:
            flush()
            heading = match.group(2).strip()
        else:
            buffer.append(line)
    flush()

    if not sections:
        raise DocumentParseError("This Markdown file has no text.")
    return sections


def _parse_text(data: bytes) -> list[Section]:
    text = _clean(_decode_text(data))
    if not text:
        raise DocumentParseError("This text file is empty.")
    return [Section(text=text)]


def parse_document(data: bytes, kind: DocumentKind) -> list[Section]:
    """Blocking; run it in a worker thread from async code."""
    if kind == "pdf":
        return _parse_pdf(data)
    if kind == "docx":
        return _parse_docx(data)
    if kind == "md":
        return _parse_markdown(data)
    return _parse_text(data)
